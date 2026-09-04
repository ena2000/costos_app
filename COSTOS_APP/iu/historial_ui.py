# iu/historial_ui.py
import tkinter as tk
from tkinter import ttk, messagebox, filedialog
from models.database import obtener_conexion
from models.historial import consultar_horas_desglose, eliminar_orden
import pandas as pd
from datetime import datetime


class HistorialUI:
    def __init__(self, parent):
        self.root = parent
        self.root.columnconfigure(0, weight=1)
        self.root.rowconfigure(2, weight=1)
        self.ultimo_desglose = None

        # --- FILTRO POR FECHAS ---
        filtro_frame = ttk.LabelFrame(self.root, text=" Filtrar horas por fechas ", padding=10)
        filtro_frame.grid(row=0, column=0, sticky="ew", padx=10, pady=5)

        ttk.Label(filtro_frame, text="Desde (DD/MM/AAAA):").pack(side="left", padx=(0, 5))
        self.entry_desde = ttk.Entry(filtro_frame, width=12)
        self.entry_desde.pack(side="left", padx=5)

        ttk.Label(filtro_frame, text="Hasta (DD/MM/AAAA):").pack(side="left", padx=(10, 5))
        self.entry_hasta = ttk.Entry(filtro_frame, width=12)
        self.entry_hasta.pack(side="left", padx=5)

        # Por defecto: mes actual
        hoy = datetime.now()
        self.entry_desde.insert(0, f"01/{hoy.strftime('%m/%Y')}")
        self.entry_hasta.insert(0, hoy.strftime("%d/%m/%Y"))

        ttk.Button(filtro_frame, text="Aplicar filtro", command=self.cargar_datos).pack(side="left", padx=8)
        ttk.Button(filtro_frame, text="Ver todo", command=self._ver_todo).pack(side="left", padx=2)
        ttk.Button(filtro_frame, text="Mes actual", command=self._mes_actual).pack(side="left", padx=2)

        # --- RESUMEN TOTALES ---
        self.resumen_frame = ttk.LabelFrame(self.root, text=" Resumen de horas (rango seleccionado) ", padding=10)
        self.resumen_frame.grid(row=1, column=0, sticky="ew", padx=10, pady=5)

        self.lbl_horas = ttk.Label(self.resumen_frame, text="Horas totales: 0.00", font=("Segoe UI", 11, "bold"))
        self.lbl_horas.pack(side="left", padx=15)
        self.lbl_maq = ttk.Label(self.resumen_frame, text="Hrs máquina: 0.00")
        self.lbl_maq.pack(side="left", padx=15)
        self.lbl_mo = ttk.Label(self.resumen_frame, text="Hrs mano obra: 0.00")
        self.lbl_mo.pack(side="left", padx=15)

        ttk.Button(self.resumen_frame, text="Exportar desglose Excel", command=self.exportar_excel).pack(side="right", padx=5)
        ttk.Button(self.resumen_frame, text="Actualizar", command=self.cargar_datos).pack(side="right", padx=5)

        # --- NOTEBOOK: Órdenes / Mes / Máquina / Operario ---
        self.notebook = ttk.Notebook(self.root)
        self.notebook.grid(row=2, column=0, sticky="nsew", padx=10, pady=5)

        self.frame_ordenes = ttk.Frame(self.notebook)
        self.frame_mes = ttk.Frame(self.notebook)
        self.frame_maquina = ttk.Frame(self.notebook)
        self.frame_operario = ttk.Frame(self.notebook)

        self.notebook.add(self.frame_ordenes, text="Órdenes")
        self.notebook.add(self.frame_mes, text="Por mes")
        self.notebook.add(self.frame_maquina, text="Por máquina")
        self.notebook.add(self.frame_operario, text="Por operario")

        self._crear_tabla_ordenes()
        self._crear_tabla_mes()
        self._crear_tabla_maquina()
        self._crear_tabla_operario()

        ttk.Button(self.root, text="Eliminar registro seleccionado", command=self.eliminar_registro).grid(
            row=3, column=0, pady=10
        )

        self.cargar_datos()

    def _crear_tabla_ordenes(self):
        self.frame_ordenes.columnconfigure(0, weight=1)
        self.frame_ordenes.rowconfigure(0, weight=1)
        columnas = ("id", "fecha", "cliente", "descripcion", "h_maq", "h_mo", "horas", "total")
        self.tree = ttk.Treeview(self.frame_ordenes, columns=columnas, show="headings")
        headings = {
            "id": ("ID", 40),
            "fecha": ("F. Inicio", 90),
            "cliente": ("Cliente", 140),
            "descripcion": ("Trabajo", 200),
            "h_maq": ("Hrs Máq.", 80),
            "h_mo": ("Hrs M.O.", 80),
            "horas": ("Hrs Tot.", 80),
            "total": ("Total $", 90),
        }
        for col, (titulo, ancho) in headings.items():
            self.tree.heading(col, text=titulo)
            self.tree.column(col, width=ancho, anchor="center" if col != "cliente" and col != "descripcion" else "w")
        self.tree.grid(row=0, column=0, sticky="nsew")
        scroll = ttk.Scrollbar(self.frame_ordenes, orient="vertical", command=self.tree.yview)
        self.tree.configure(yscrollcommand=scroll.set)
        scroll.grid(row=0, column=1, sticky="ns")

    def _crear_tabla_mes(self):
        self.frame_mes.columnconfigure(0, weight=1)
        self.frame_mes.rowconfigure(0, weight=1)
        cols = ("mes", "ordenes", "h_maq", "h_mo", "h_tot")
        self.tree_mes = ttk.Treeview(self.frame_mes, columns=cols, show="headings")
        for col, titulo, w in (
            ("mes", "Mes", 160),
            ("ordenes", "Órdenes", 80),
            ("h_maq", "Hrs Máquina", 110),
            ("h_mo", "Hrs M.O.", 110),
            ("h_tot", "Hrs Totales", 110),
        ):
            self.tree_mes.heading(col, text=titulo)
            self.tree_mes.column(col, width=w, anchor="center")
        self.tree_mes.grid(row=0, column=0, sticky="nsew")
        scroll = ttk.Scrollbar(self.frame_mes, orient="vertical", command=self.tree_mes.yview)
        self.tree_mes.configure(yscrollcommand=scroll.set)
        scroll.grid(row=0, column=1, sticky="ns")

    def _crear_tabla_maquina(self):
        self.frame_maquina.columnconfigure(0, weight=1)
        self.frame_maquina.rowconfigure(0, weight=1)
        cols = ("maquina", "horas", "operarios")
        self.tree_maq = ttk.Treeview(self.frame_maquina, columns=cols, show="headings")
        for col, titulo, w in (
            ("maquina", "Máquina", 220),
            ("horas", "Horas", 100),
            ("operarios", "Desglose por operario / diseñador", 420),
        ):
            self.tree_maq.heading(col, text=titulo)
            self.tree_maq.column(col, width=w, anchor="w" if col != "horas" else "center")
        self.tree_maq.grid(row=0, column=0, sticky="nsew")
        scroll = ttk.Scrollbar(self.frame_maquina, orient="vertical", command=self.tree_maq.yview)
        self.tree_maq.configure(yscrollcommand=scroll.set)
        scroll.grid(row=0, column=1, sticky="ns")

    def _crear_tabla_operario(self):
        self.frame_operario.columnconfigure(0, weight=1)
        self.frame_operario.rowconfigure(0, weight=1)
        cols = ("operario", "tipo", "horas", "detalle")
        self.tree_ope = ttk.Treeview(self.frame_operario, columns=cols, show="headings")
        for col, titulo, w in (
            ("operario", "Operario / Diseñador", 200),
            ("tipo", "Tipo", 100),
            ("horas", "Horas", 100),
            ("detalle", "Desglose (máquina / actividad)", 380),
        ):
            self.tree_ope.heading(col, text=titulo)
            self.tree_ope.column(col, width=w, anchor="w" if col not in ("horas", "tipo") else "center")
        self.tree_ope.grid(row=0, column=0, sticky="nsew")
        scroll = ttk.Scrollbar(self.frame_operario, orient="vertical", command=self.tree_ope.yview)
        self.tree_ope.configure(yscrollcommand=scroll.set)
        scroll.grid(row=0, column=1, sticky="ns")

    def _mes_actual(self):
        hoy = datetime.now()
        self.entry_desde.delete(0, tk.END)
        self.entry_desde.insert(0, f"01/{hoy.strftime('%m/%Y')}")
        self.entry_hasta.delete(0, tk.END)
        self.entry_hasta.insert(0, hoy.strftime("%d/%m/%Y"))
        self.cargar_datos()

    def _ver_todo(self):
        self.entry_desde.delete(0, tk.END)
        self.entry_hasta.delete(0, tk.END)
        self.cargar_datos()

    def _limpiar_tree(self, tree):
        for item in tree.get_children():
            tree.delete(item)

    def cargar_datos(self):
        desde = self.entry_desde.get().strip() or None
        hasta = self.entry_hasta.get().strip() or None

        try:
            desglose = consultar_horas_desglose(desde, hasta)
            self.ultimo_desglose = desglose
        except Exception as e:
            messagebox.showerror("Error de carga", f"No se pudo leer el historial: {e}")
            return

        self._limpiar_tree(self.tree)
        self._limpiar_tree(self.tree_mes)
        self._limpiar_tree(self.tree_maq)
        self._limpiar_tree(self.tree_ope)

        for o in desglose["ordenes"]:
            self.tree.insert("", "end", values=(
                o["id"],
                o["fecha_inicio"],
                o["cliente"],
                o["descripcion"],
                f"{o['horas_maquina']:.2f}",
                f"{o['horas_mano_obra']:.2f}",
                f"{o['horas_totales']:.2f}",
                f"$ {o['costo_total']:.2f}",
            ))

        for m in desglose["por_mes"]:
            self.tree_mes.insert("", "end", values=(
                m.get("etiqueta", ""),
                m.get("ordenes", 0),
                f"{m.get('horas_maquina', 0):.2f}",
                f"{m.get('horas_mo', 0):.2f}",
                f"{m.get('horas_totales', 0):.2f}",
            ))

        for m in desglose["por_maquina"]:
            ops = m.get("operarios") or {}
            detalle_ops = ", ".join(f"{nombre}: {hrs:.2f}h" for nombre, hrs in sorted(ops.items()))
            self.tree_maq.insert("", "end", values=(
                m["maquina"],
                f"{m['horas']:.2f}",
                detalle_ops or "-",
            ))

        for o in desglose["por_operario"]:
            conceptos = o.get("conceptos") or {}
            detalle = ", ".join(f"{c}: {h:.2f}h" for c, h in sorted(conceptos.items()))
            self.tree_ope.insert("", "end", values=(
                o["operario"],
                o["tipo"],
                f"{o['horas']:.2f}",
                detalle or "-",
            ))

        self.lbl_horas.config(text=f"Horas totales: {desglose['horas_totales']:.2f}")
        self.lbl_maq.config(text=f"Hrs máquina: {desglose['horas_maquina']:.2f}")
        self.lbl_mo.config(text=f"Hrs mano obra: {desglose['horas_mano_obra']:.2f}")

    def exportar_excel(self):
        """Exporta horas totales + desglose por mes, máquina y operario."""
        try:
            desde = self.entry_desde.get().strip() or None
            hasta = self.entry_hasta.get().strip() or None
            desglose = self.ultimo_desglose or consultar_horas_desglose(desde, hasta)

            if not desglose["ordenes"] and not desglose["por_maquina"] and not desglose["por_operario"]:
                messagebox.showwarning("Sin datos", "No hay registros en el rango seleccionado.")
                return

            ruta = filedialog.asksaveasfilename(
                defaultextension=".xlsx",
                filetypes=[("Archivos Excel", "*.xlsx")],
                initialfile=f"Horas_Desglose_{datetime.now().strftime('%Y_%m_%d')}.xlsx",
            )
            if not ruta:
                return

            from openpyxl.styles import Font, PatternFill, Alignment

            df_resumen = pd.DataFrame([
                {"Concepto": "Horas totales", "Valor": round(desglose["horas_totales"], 2)},
                {"Concepto": "Horas máquina", "Valor": round(desglose["horas_maquina"], 2)},
                {"Concepto": "Horas mano de obra", "Valor": round(desglose["horas_mano_obra"], 2)},
                {"Concepto": "Órdenes en rango", "Valor": len(desglose["ordenes"])},
                {"Concepto": "Fecha desde", "Valor": desde or "(todas)"},
                {"Concepto": "Fecha hasta", "Valor": hasta or "(todas)"},
            ])

            df_mes = pd.DataFrame([
                {
                    "Mes": m.get("etiqueta"),
                    "Órdenes": m.get("ordenes"),
                    "Hrs Máquina": round(m.get("horas_maquina", 0), 2),
                    "Hrs M.O.": round(m.get("horas_mo", 0), 2),
                    "Hrs Totales": round(m.get("horas_totales", 0), 2),
                }
                for m in desglose["por_mes"]
            ])

            filas_maq = []
            for m in desglose["por_maquina"]:
                ops = m.get("operarios") or {}
                if ops:
                    for nombre, hrs in ops.items():
                        filas_maq.append({
                            "Máquina": m["maquina"],
                            "Operario / Diseñador": nombre,
                            "Horas": round(hrs, 2),
                            "Total máquina": round(m["horas"], 2),
                        })
                else:
                    filas_maq.append({
                        "Máquina": m["maquina"],
                        "Operario / Diseñador": "-",
                        "Horas": round(m["horas"], 2),
                        "Total máquina": round(m["horas"], 2),
                    })
            df_maq = pd.DataFrame(filas_maq)

            filas_ope = []
            for o in desglose["por_operario"]:
                conceptos = o.get("conceptos") or {}
                if conceptos:
                    for concepto, hrs in conceptos.items():
                        filas_ope.append({
                            "Operario / Diseñador": o["operario"],
                            "Tipo": o["tipo"],
                            "Concepto (máquina/actividad)": concepto,
                            "Horas": round(hrs, 2),
                            "Total persona": round(o["horas"], 2),
                        })
                else:
                    filas_ope.append({
                        "Operario / Diseñador": o["operario"],
                        "Tipo": o["tipo"],
                        "Concepto (máquina/actividad)": "-",
                        "Horas": round(o["horas"], 2),
                        "Total persona": round(o["horas"], 2),
                    })
            df_ope = pd.DataFrame(filas_ope)

            df_ordenes = pd.DataFrame([
                {
                    "ID": o["id"],
                    "F. Inicio": o["fecha_inicio"],
                    "F. Fin": o["fecha_fin"],
                    "Cliente": o["cliente"],
                    "Trabajo": o["descripcion"],
                    "Hrs Máquina": round(o["horas_maquina"], 2),
                    "Hrs M.O.": round(o["horas_mano_obra"], 2),
                    "Hrs Totales": round(o["horas_totales"], 2),
                    "Total $": round(o["costo_total"], 2),
                }
                for o in desglose["ordenes"]
            ])

            # Historial completo (costos) del rango
            conn = obtener_conexion()
            df_full = pd.read_sql_query("SELECT * FROM historial_ordenes ORDER BY id DESC", conn)
            conn.close()
            ids_rango = {o["id"] for o in desglose["ordenes"]}
            if not df_full.empty and ids_rango:
                df_full = df_full[df_full["id"].isin(ids_rango)]

            with pd.ExcelWriter(ruta, engine="openpyxl") as writer:
                df_resumen.to_excel(writer, index=False, sheet_name="Resumen Totales")
                if not df_mes.empty:
                    df_mes.to_excel(writer, index=False, sheet_name="Por Mes")
                if not df_maq.empty:
                    df_maq.to_excel(writer, index=False, sheet_name="Por Maquina")
                if not df_ope.empty:
                    df_ope.to_excel(writer, index=False, sheet_name="Por Operario")
                if not df_ordenes.empty:
                    df_ordenes.to_excel(writer, index=False, sheet_name="Ordenes")
                if not df_full.empty:
                    df_full.to_excel(writer, index=False, sheet_name="Historial Completo")

                header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
                header_font = Font(bold=True, color="FFFFFF")
                center_align = Alignment(horizontal="center", vertical="center")

                for sheet_name in writer.sheets:
                    ws = writer.sheets[sheet_name]
                    for cell in ws[1]:
                        cell.fill = header_fill
                        cell.font = header_font
                        cell.alignment = center_align
                    for column in ws.columns:
                        max_length = 0
                        column_letter = column[0].column_letter
                        for cell in column:
                            try:
                                max_length = max(max_length, len(str(cell.value)))
                            except Exception:
                                pass
                        ws.column_dimensions[column_letter].width = min(max_length + 3, 45)

            messagebox.showinfo("Éxito", f"Desglose exportado:\n{ruta}")

        except Exception as e:
            messagebox.showerror("Error de exportación", f"No se pudo generar el Excel: {e}")

    def eliminar_registro(self):
        selected = self.tree.selection()
        if not selected:
            messagebox.showwarning("Atención", "Selecciona una fila en la pestaña Órdenes.")
            return

        if messagebox.askyesno("Confirmar", "¿Seguro que deseas borrar este registro del historial?"):
            item_id = self.tree.item(selected[0])["values"][0]
            try:
                eliminar_orden(item_id)
                self.cargar_datos()
                messagebox.showinfo("Éxito", "Registro eliminado.")
            except Exception as e:
                messagebox.showerror("Error", f"No se pudo eliminar: {e}")
